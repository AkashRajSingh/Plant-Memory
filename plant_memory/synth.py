"""Synthetic SAP PM export generator for a rice processing unit.

Produces an Excel file shaped like the plant's real export (one row per
material issue, header fields repeated per line, messy short texts) plus a
hidden ground-truth file used only by the evaluation harness.

Planted patterns the product should surface:
- Dryer elevator is a bad actor: ~2x failure rate, dominated by misalignment.
- V-belt C-106 is consumed across several parboiling elevators.
- Husker rubber rolls are the biggest recurring spend in milling.
- From July 2024, the Hodi conveyor starts showing gearbox oil leaks caused by
  a choked breather: a failure mode and cause the taxonomy does not know yet
  (ontology drift).
"""
from __future__ import annotations

import csv
import math
import random
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

from .catalog import DEFAULT_ASSETS as ASSETS, DEFAULT_MATERIALS as MATERIALS, DEFAULT_PROCESS_AREAS as PROCESS_AREAS, materials_for

# --- failure scenarios per asset class ------------------------------------
# (component, failure_mode, weight, symptoms, causes{cause: w}, actions, extra components, qty range, downtime hrs)
S = dict
SCENARIOS = {
    "bucket_elevator": [
        S(c="bearing", fm="bearing_wear", w=3, sym=["abnormal_noise", "vibration", "overheating"],
          causes={"inadequate_lubrication": 4, "misalignment": 3, "dust_ingress": 2, "normal_wear": 2, "material_defect": 1},
          act=["replace"], extra=["bearing_sleeve", "v_belt"], qty=(2, 2), dt=(2, 6)),
        S(c="v_belt", fm="belt_slippage", w=2, sym=["belt_slip", "reduced_output"],
          causes={"improper_tensioning": 5, "normal_wear": 3, "overload": 2}, act=["replace", "adjust"], extra=[], qty=(2, 3), dt=(0.5, 2)),
        S(c="sprocket", fm="sprocket_tooth_wear", w=2, sym=["abnormal_noise", "vibration"],
          causes={"misalignment": 4, "normal_wear": 4, "inadequate_lubrication": 2}, act=["replace"], extra=[], qty=(1, 2), dt=(1.5, 4)),
        S(c="bucket", fm="bucket_damage", w=1, sym=["stopped", "abnormal_noise"],
          causes={"foreign_material": 6, "overload": 4}, act=["replace"], extra=[], qty=(6, 20), dt=(2, 5)),
        S(c="elevator_belt", fm="belt_wear", w=1, sym=["belt_slip", "stopped"],
          causes={"normal_wear": 5, "misalignment": 3, "overload": 2}, act=["replace"], extra=[], qty=(8, 30), dt=(6, 14)),
        S(c="chain_coupling", fm="coupling_failure", w=1, sym=["stopped", "vibration"],
          causes={"misalignment": 6, "normal_wear": 2, "improper_installation": 2}, act=["replace"], extra=[], qty=(1, 2), dt=(2, 5)),
        S(c="shaft", fm="shaft_fracture", w=0.4, sym=["stopped"],
          causes={"overload": 5, "misalignment": 3, "material_defect": 2}, act=["weld_repair"], extra=[], qty=(1, 2), dt=(8, 20)),
    ],
    "belt_conveyor": [
        S(c="conveyor_belt", fm="belt_wear", w=2, sym=["belt_slip", "stopped"],
          causes={"normal_wear": 5, "misalignment": 3, "foreign_material": 2}, act=["replace"], extra=["sprocket"], qty=(4, 12), dt=(4, 10)),
        S(c="plummer_block", fm="bearing_wear", w=2, sym=["abnormal_noise", "overheating"],
          causes={"inadequate_lubrication": 5, "dust_ingress": 3, "normal_wear": 2}, act=["replace"], extra=["bearing", "bearing_sleeve"], qty=(1, 2), dt=(2, 5)),
        S(c="shaft", fm="shaft_fracture", w=0.6, sym=["stopped"],
          causes={"overload": 6, "material_defect": 4}, act=["weld_repair"], extra=[], qty=(1, 2), dt=(6, 16)),
    ],
    "husker": [
        S(c="rubber_roll", fm="rubber_roll_wear", w=6, sym=["reduced_output", "quality_issue"],
          causes={"normal_wear": 6, "operator_error": 3, "moisture_ingress": 1}, act=["replace"], extra=[], qty=(1, 2), dt=(0.5, 1.5)),
        S(c="bearing", fm="bearing_wear", w=2, sym=["abnormal_noise", "overheating", "vibration"],
          causes={"dust_ingress": 5, "inadequate_lubrication": 3, "normal_wear": 2}, act=["replace"], extra=[], qty=(2, 3), dt=(2, 4)),
        S(c="v_belt", fm="belt_wear", w=1, sym=["belt_slip"],
          causes={"normal_wear": 6, "improper_tensioning": 4}, act=["replace"], extra=[], qty=(2, 4), dt=(0.5, 1.5)),
    ],
    "whitener": [
        S(c="whitening_stone", fm="stone_wear", w=4, sym=["quality_issue", "reduced_output"],
          causes={"normal_wear": 8, "foreign_material": 2}, act=["replace"], extra=[], qty=(1, 1), dt=(2, 4)),
        S(c="screen_deck", fm="screen_tear", w=1, sym=["quality_issue"],
          causes={"foreign_material": 6, "normal_wear": 4}, act=["replace"], extra=[], qty=(1, 2), dt=(1, 3)),
        S(c="bearing", fm="bearing_wear", w=1, sym=["abnormal_noise", "overheating"],
          causes={"dust_ingress": 5, "inadequate_lubrication": 5}, act=["replace"], extra=[], qty=(2, 2), dt=(2, 5)),
    ],
    "polisher": [
        S(c="polishing_screen", fm="screen_tear", w=3, sym=["quality_issue"],
          causes={"normal_wear": 6, "foreign_material": 4}, act=["replace"], extra=[], qty=(1, 2), dt=(1, 3)),
        S(c="bearing", fm="bearing_seizure", w=1, sym=["overheating", "stopped"],
          causes={"inadequate_lubrication": 6, "moisture_ingress": 4}, act=["replace"], extra=[], qty=(2, 2), dt=(3, 6)),
    ],
    "color_sorter": [
        S(c="ejector_valve", fm="valve_sticking", w=4, sym=["quality_issue"],
          causes={"dust_ingress": 6, "moisture_ingress": 2, "normal_wear": 2}, act=["replace", "clean"], extra=[], qty=(2, 8), dt=(0.5, 2)),
        S(c="electrical_fitting", fm="electrical_fault", w=1, sym=["stopped", "tripping"],
          causes={"electrical_supply": 7, "dust_ingress": 3}, act=["repair"], extra=[], qty=(1, 3), dt=(1, 3)),
    ],
    "vibro_screen": [
        S(c="screen_deck", fm="screen_tear", w=3, sym=["quality_issue", "reduced_output"],
          causes={"foreign_material": 5, "normal_wear": 5}, act=["replace"], extra=["cleaning_cup"], qty=(2, 4), dt=(3, 6)),
        S(c="bearing", fm="bearing_wear", w=1, sym=["vibration", "abnormal_noise"],
          causes={"dust_ingress": 6, "normal_wear": 4}, act=["replace"], extra=[], qty=(2, 2), dt=(2, 4)),
    ],
    "pump": [
        S(c="mechanical_seal", fm="seal_leakage", w=3, sym=["leakage"],
          causes={"cavitation": 4, "misalignment": 3, "normal_wear": 3}, act=["replace", "refurbish"], extra=["coupling"], qty=(1, 1), dt=(2, 6)),
        S(c="impeller", fm="impeller_erosion", w=1, sym=["reduced_output", "vibration"],
          causes={"cavitation": 7, "normal_wear": 3}, act=["refurbish", "replace"], extra=["bearing"], qty=(1, 1), dt=(4, 10)),
        S(c="motor", fm="motor_overheating", w=1, sym=["tripping", "overheating"],
          causes={"electrical_supply": 5, "overload": 5}, act=["repair"], extra=["bearing"], qty=(2, 2), dt=(3, 8)),
    ],
    "dryer": [
        S(c="temperature_sensor", fm="sensor_failure", w=3, sym=["quality_issue", "overheating"],
          causes={"moisture_ingress": 6, "normal_wear": 4}, act=["replace"], extra=[], qty=(1, 1), dt=(1, 3)),
        S(c="bearing", fm="bearing_wear", w=1, sym=["abnormal_noise"],
          causes={"moisture_ingress": 5, "inadequate_lubrication": 5}, act=["replace"], extra=[], qty=(2, 2), dt=(3, 6)),
    ],
    "packing_machine": [
        S(c="bearing", fm="bearing_wear", w=2, sym=["abnormal_noise", "vibration"],
          causes={"dust_ingress": 5, "normal_wear": 5}, act=["replace"], extra=[], qty=(2, 2), dt=(1, 3)),
        S(c="electrical_fitting", fm="electrical_fault", w=2, sym=["stopped", "tripping"],
          causes={"electrical_supply": 6, "dust_ingress": 4}, act=["repair"], extra=[], qty=(1, 2), dt=(1, 4)),
    ],
}

# Failures per asset per month (Poisson mean)
CLASS_RATE = {"bucket_elevator": 1.3, "belt_conveyor": 0.9, "husker": 2.0, "whitener": 0.9, "polisher": 0.7,
              "color_sorter": 0.8, "vibro_screen": 0.5, "pump": 0.5, "dryer": 0.4, "packing_machine": 0.6}
ASSET_RATE_MULT = {"PARB-DRY-ELV-01": 2.2}
ASSET_CAUSE_BIAS = {"PARB-DRY-ELV-01": {"misalignment": 4.0}}

# --- how people write things (deliberately messier than the taxonomy) --------
COMP_TEXT = {
    "bearing": ["Bearing Change", "brg change", "Brg chng", "bearing replaced", "Air Bearing Change", "brg replacement"],
    "plummer_block": ["Plmbr Blck Change", "plummer block chng", "Plumber block change"],
    "v_belt": ["v belt change", "V-Belt chng", "belt change", "vbelt replacement"],
    "elevator_belt": ["elevator belt change", "Elev belt replacement", "bucket belt chng"],
    "conveyor_belt": ["Conveyor Belt Change", "conv belt replaced", "belt joint & change"],
    "sprocket": ["sprocket Changed", "Sprocket change", "sprkt replacement"],
    "chain_coupling": ["Chain Coupling Change", "chain cplng chng"],
    "coupling": ["coupling change", "cplng replaced"],
    "shaft": ["SHAFT BROKEN(REPAIR)", "shaft repair welding", "shaft broke - weld"],
    "bucket": ["bucket change", "bkt replacement", "buckets fitted"],
    "screen_deck": ["deck change", "screen change", "sieve replacement"],
    "rubber_roll": ["rubber roll change", "rbr roll chng", "roll replacement", "rubber roll replaced"],
    "whitening_stone": ["stone change", "emery stone chng", "whitening stone replaced"],
    "polishing_screen": ["screen change", "polisher screen chng"],
    "mechanical_seal": ["mech seal change", "seal leak rectified", "gland leakage"],
    "impeller": ["impeller change", "impeller replaced"],
    "ejector_valve": ["ejector chng", "ejector valve replacement", "nozzle cleaning"],
    "temperature_sensor": ["RTD change", "temp sensor replaced", "PT-100 chng"],
    "motor": ["motor rewinding", "mtr repair", "motor heating problem"],
    "electrical_fitting": ["electrical work", "panel wiring repair", "supply problem"],
    "gearbox": ["gearbox oil leak", "GB oil leakage", "gear box leaking"],
}
SYMPTOM_TEXT = {"abnormal_noise": ["Noise", "Abnormal sound", "Noisy"], "vibration": ["Vibration high", "Vibration"],
                "overheating": ["Heating", "Temperature high"], "belt_slip": ["Belt slipping", "Slip"],
                "stopped": ["Breakdown", "Stopped", "B/D"], "leakage": ["Leakage", "Leak"],
                "reduced_output": ["Low output", "Capacity drop", "Less husking"],
                "quality_issue": ["Broken rice high", "Quality issue", "Colour reject high"],
                "tripping": ["Tripping", "Motor trip"]}
RCA_TEXT = {"misalignment": ["Misalignment", "Pulley alignment out", "Mis-alignment"],
            "inadequate_lubrication": ["Lubrication failure", "Lack of greasing", "No grease"],
            "dust_ingress": ["Dust ingress", "Dust inside housing"],
            "moisture_ingress": ["Moisture ingress", "Steam/water ingress"],
            "normal_wear": ["Wear & tear", "Life over", "Normal wear"],
            "improper_tensioning": ["Belt tension improper", "Loose belt"],
            "overload": ["Overload", "Choke-up"],
            "foreign_material": ["Foreign material", "Stone in feed"],
            "material_defect": ["Poor quality spare", "Local make failed"],
            "improper_installation": ["Improper fitment", "Not tightened properly"],
            "operator_error": ["Wrong roll gap setting", "Operator setting"],
            "cavitation": ["Cavitation", "Dry running"],
            "electrical_supply": ["Voltage fluctuation", "Single phasing"],
            "breather_choked": ["Breather choked"]}
CATEGORY_TEXT = {"misalignment": "Mechanical", "inadequate_lubrication": "Maintenance", "dust_ingress": "Environment",
                 "moisture_ingress": "Process", "normal_wear": "Mechanical", "improper_tensioning": "Maintenance",
                 "overload": "Operational", "foreign_material": "Process", "material_defect": "Material quality",
                 "improper_installation": "Maintenance", "operator_error": "Operational", "cavitation": "Process",
                 "electrical_supply": "Electrical", "breather_choked": "Maintenance"}
# Technician remarks: free text with cause cues, some phrased outside the taxonomy synonyms
REMARK_TEXT = {
    "misalignment": ["head pulley alignment found off ~3mm, corrected", "shaft out of line with motor", "alignment sahi kiya, pulley shifted",
                     "head and tail pulley not matching"],
    "inadequate_lubrication": ["no grease in brg housing", "grease line choked, bearing dry", "grease nahi tha, brg jam"],
    "dust_ingress": ["husk dust accumulation inside", "seal gone, dust entered", "bran dust packed inside"],
    "moisture_ingress": ["water ingress from steam line", "condensate dripping on unit", "wet paddy sticking"],
    "normal_wear": ["worn out after long run", "life over, normal replacement", "running since last season"],
    "improper_tensioning": ["belt loose, tension adjusted", "tensioner bolt slack"],
    "overload": ["choke-up in boot, overload", "excess feed from hodi, jam", "feed gate fully open"],
    "foreign_material": ["stone found in boot", "metal piece found in bucket line", "iron piece came with paddy"],
    "material_defect": ["local make brg failed early", "new spare was defective", "poor quality belt, cracked in 2 weeks"],
    "improper_installation": ["sleeve not tightened properly last time", "wrong fitment by contractor"],
    "operator_error": ["roll gap set wrong by operator", "operator changed setting on shift"],
    "cavitation": ["pump running dry, low sump level", "suction strainer choked"],
    "electrical_supply": ["voltage fluctuation from grid", "single phasing on R phase"],
    "breather_choked": ["breather choked with husk dust, pressure built up", "breather blocked, oil pushed out of seal"],
}
ACTION_TEXT = {"replace": ["Replaced", "Changed", "New fitted"], "adjust": ["Tension adjusted", "Tightened"],
               "weld_repair": ["Welding done", "Weld repair"], "repair": ["Repaired", "Rectified"],
               "refurbish": ["Refurbished", "Overhauled"], "clean": ["Cleaned"], "lubricate": ["Greased"],
               "modify": ["Fabricated", "Modification done"], "realign": ["Aligned"]}
TECHNICIANS = ["T1043", "T1057", "T1102", "T1118", "T1125", "T1201"]
COST_CENTERS = ["121FPTOR", "121UBTEC", "121PARBM", "121MILLM"]

COLUMNS = ["Order", "Posting Date", "Work Description", "Material description", "Cost (In Rs)", "Quantity",
           "Base Unit of Measure", "G/L Account", "Cost Center", "Planning Department", "Work Type", "Movement type",
           "Material", "Storage location", "Functional Location", "Symptom", "Root Cause", "Cause Category",
           "Action Taken", "Downtime (Hrs)", "Breakdown", "Technician Remarks", "Technician"]


@dataclass
class TruthEvent:
    order: str
    date: date
    asset_id: str | None
    kind: str  # failure | lubrication | modification | reservation
    component: str | None = None
    failure_mode: str | None = None
    root_cause: str | None = None
    action: str | None = None
    symptom: str | None = None
    breakdown: bool = False
    lines: list = field(default_factory=list)


def _wchoice(rng: random.Random, weights: dict):
    keys = list(weights)
    return rng.choices(keys, weights=[weights[k] for k in keys])[0]


def _poisson(rng: random.Random, lam: float) -> int:
    l, k, p = math.exp(-lam), 0, 1.0
    while True:
        p *= rng.random()
        if p <= l:
            return k
        k += 1


class _OrderNumbers:
    def __init__(self):
        self.n = {"3": 230100, "2": 150500, "4": 1500, "8": 10750}

    def next(self, kind: str) -> str:
        self.n[kind] += 1
        prefix = {"3": "3000", "2": "2000", "4": "4000", "8": "8000"}[kind]
        return f"{prefix}{self.n[kind]:07d}"


def _maybe_upper(rng, s):
    return s.upper() if rng.random() < 0.08 else s


def _description(rng, asset, scen_component, symptom, kind) -> str:
    area = PROCESS_AREAS[asset["area"]]
    prefix = rng.choice(area["text_prefixes"])
    alias = rng.choice(asset["aliases"])
    if kind == "lubrication":
        txt = f"{prefix}{alias} {rng.choice(['greasing', 'lubrication round', 'PM greasing'])}"
    elif kind == "generic":
        txt = f"{prefix}{alias} {rng.choice(['maintenance', 'maint work', 'PM work', 'problem attended'])}"
    elif kind == "adjust":
        txt = f"{prefix}{alias} {rng.choice(['belt tightening', 'v belt tension adjust', 'belt slip attended'])}"
    else:
        comp_txt = rng.choice(COMP_TEXT.get(scen_component, [scen_component.replace('_', ' ') + ' change']))
        if symptom and rng.random() < 0.25:
            sym = rng.choice({"abnormal_noise": ["noisy", "noise"], "stopped": ["brkdn", "B/D"], "leakage": ["leak"],
                              "vibration": ["vibration"], "belt_slip": ["slip"]}.get(symptom, [""]))
            comp_txt = f"{sym} {comp_txt}".strip()
        txt = f"{prefix}{alias} {comp_txt}"
    return _maybe_upper(rng, txt)[:40]


def _lines_for(rng, components, asset_class, qty_range):
    lines = []
    for i, comp in enumerate(components):
        mats = materials_for(comp, asset_class)
        if not mats:
            continue
        n_mats = 1 if i else rng.choice([1, 1, 2]) if len(mats) > 1 else 1
        for code in rng.sample(mats, min(n_mats, len(mats))):
            q = rng.randint(*qty_range) if i == 0 else rng.randint(1, 2)
            lines.append((code, q))
    return lines


def generate(out_xlsx: Path, truth_csv: Path, start: date = date(2023, 10, 1), months: int = 12, seed: int = 7,
             rate_scale: float = 1.6) -> dict:
    rng = random.Random(seed)
    orders = _OrderNumbers()
    events: list[TruthEvent] = []

    for m in range(months):
        month_start = date(start.year + (start.month - 1 + m) // 12, (start.month - 1 + m) % 12 + 1, 1)

        def rand_day():
            return month_start + timedelta(days=rng.randint(0, 27))

        for aid, asset in ASSETS.items():
            cls = asset["class"]
            # failures
            if cls in SCENARIOS:
                lam = CLASS_RATE[cls] * ASSET_RATE_MULT.get(aid, 1.0) * rate_scale
                for _ in range(_poisson(rng, lam)):
                    scen = rng.choices(SCENARIOS[cls], weights=[s["w"] for s in SCENARIOS[cls]])[0]
                    causes = dict(scen["causes"])
                    for c, mult in ASSET_CAUSE_BIAS.get(aid, {}).items():
                        if c in causes:
                            causes[c] *= mult
                    cause = _wchoice(rng, causes)
                    breakdown = scen["fm"] in {"shaft_fracture", "coupling_failure", "bucket_damage", "bearing_seizure",
                                               "motor_overheating", "electrical_fault"} or rng.random() < 0.45
                    action = scen["act"][0] if len(scen["act"]) == 1 or rng.random() < 0.7 else scen["act"][1]
                    kind_code = "4" if action == "refurbish" else ("2" if breakdown and rng.random() < 0.7 else "3")
                    ev = TruthEvent(orders.next(kind_code), rand_day(), aid, "failure", scen["c"], scen["fm"], cause,
                                    action, rng.choice(scen["sym"]), breakdown)
                    comps = [scen["c"]] + [c for c in scen["extra"] if rng.random() < 0.6]
                    ev.lines = _lines_for(rng, comps, cls, scen["qty"]) if action != "adjust" else []
                    ev._dt = round(rng.uniform(*scen["dt"]), 1)
                    events.append(ev)
            # monthly lubrication round on conveying equipment
            if cls in {"bucket_elevator", "belt_conveyor"} and rng.random() < 0.35:
                ev = TruthEvent(orders.next("3"), rand_day(), aid, "lubrication", "bearing", None, None, "lubricate")
                ev.lines = [("7010010001", rng.randint(1, 3))]
                ev._dt = 0
                events.append(ev)

        # drift: gearbox oil leaks on the Hodi conveyor from month 10 onward
        if m >= 9:
            for _ in range(rng.randint(1, 3)):
                ev = TruthEvent(orders.next("2"), rand_day(), "PARB-HODI-CNV-01", "failure", "gearbox",
                                "gearbox_oil_leakage", "breather_choked", "replace", "leakage", True)
                ev.lines = [("7010030003", 1), ("7010020002", rng.randint(5, 12))]
                ev._dt = round(rng.uniform(2, 5), 1)
                events.append(ev)

        # occasional modification jobs
        for _ in range(rng.randint(0, 2)):
            aid = rng.choice(["WH-TRL-01", "PCK-PKM-01", "PARB-HODI-CNV-01"])
            ev = TruthEvent(orders.next("8"), rand_day(), aid, "modification", None, None, None, "modify")
            ev.lines = [("5545012833", rng.choice([6, 8, 12]))] if aid == "WH-TRL-01" else [("6496032443", 1), ("6792768333", 2)]
            ev._dt = 0
            events.append(ev)

        # MB21 reservations straight to cost centre (no order)
        for _ in range(rng.randint(2, 4)):
            code = rng.choice(["5814220023", "6496033133", "6044046033", "6542025923", "7010010001"])
            ev = TruthEvent("", rand_day(), None, "reservation")
            ev.lines = [(code, rng.randint(1, 8))]
            ev._dt = 0
            events.append(ev)

    events.sort(key=lambda e: e.date)
    rows, truth = [], []
    for ev in events:
        asset = ASSETS.get(ev.asset_id) if ev.asset_id else None
        if ev.kind == "reservation":
            code, q = ev.lines[0]
            desc, cost, uom, _ = MATERIALS[code]
            rows.append({"Order": None, "Posting Date": ev.date, "Work Description": "Reservations by MB21",
                         "Material description": desc, "Cost (In Rs)": round(cost * q, 2), "Quantity": -q,
                         "Base Unit of Measure": uom, "G/L Account": 40000500, "Cost Center": rng.choice(COST_CENTERS),
                         "Planning Department": "Operations", "Work Type": "Preventive",
                         "Movement type": rng.choice([201, 201, 221]), "Material": int(code), "Storage location": "SN01"})
            truth.append({"order": "", "asset_id": "", "kind": "reservation"})
            continue

        failure = ev.kind == "failure"
        if ev.kind == "modification":
            work_type, desc_kind = "Modification", "modification"
        elif ev.kind == "lubrication":
            work_type, desc_kind = "Preventive", "lubrication"
        elif ev.action == "refurbish":
            work_type, desc_kind = "Refurbishment", "failure"
        else:
            work_type = "Corrective" if ev.breakdown and rng.random() < 0.7 else "Preventive"  # mislabels on purpose
            desc_kind = "generic" if rng.random() < 0.12 else "failure"

        if ev.kind == "modification":
            description = {"WH-TRL-01": rng.choice([f"{rng.choice([2, 3, 4])} new rice loading trolley prepare(P2R)",
                                                    "Warehouse- New hand trolley modificatn"]),
                           "PCK-PKM-01": rng.choice(["Packing- FFS 1 sensor bracket modificatn", "Packing m/c 1 panel shifting"]),
                           }.get(ev.asset_id, f"Parb- {asset['aliases'][0]} guard fabrication")[:40]
        elif ev.action == "refurbish":
            description = f"REFURB- {asset['aliases'][0]}"[:40]
        else:
            description = _description(rng, asset, ev.component, ev.symptom,
                                       "adjust" if ev.action == "adjust" and desc_kind == "failure" else desc_kind)

        header = {
            "Functional Location": ev.asset_id if rng.random() < 0.85 else None,
            "Symptom": rng.choice(SYMPTOM_TEXT[ev.symptom]) if failure and ev.symptom and rng.random() < (0.75 if ev.breakdown else 0.3) else None,
            "Root Cause": None, "Cause Category": None,
            "Action Taken": rng.choice(ACTION_TEXT[ev.action]) if ev.action and rng.random() < 0.8 else None,
            "Downtime (Hrs)": ev._dt if failure and ev.breakdown and rng.random() < 0.9 else (0 if not failure else None),
            "Breakdown": ("Yes" if ev.breakdown else "No") if rng.random() < 0.95 else None,
            "Technician Remarks": None,
            "Technician": rng.choice(TECHNICIANS),
        }
        if failure:
            if rng.random() < (0.5 if ev.breakdown else 0.2):
                header["Root Cause"] = rng.choice(RCA_TEXT[ev.root_cause])
                if rng.random() < 0.7:
                    header["Cause Category"] = CATEGORY_TEXT[ev.root_cause]
            if rng.random() < 0.55:
                header["Technician Remarks"] = rng.choice(REMARK_TEXT[ev.root_cause]) if rng.random() < 0.75 else rng.choice(
                    ["job done", "checked ok after trial", "spare taken from store", "trial run taken"])
        dept = "Electrical" if ev.component in {"electrical_fitting", "temperature_sensor", "motor"} else rng.choice(["Mechanical"] * 5 + ["Operations"])

        lines = ev.lines or []
        if not lines:  # labour-only job still appears as one line with no material
            lines = [(None, 0)]
        second_date = ev.date + timedelta(days=rng.randint(1, 4)) if rng.random() < 0.1 else None
        for i, (code, q) in enumerate(lines):
            if code:
                mdesc, ucost, uom, _ = MATERIALS[code]
                cost = round(ucost * q * rng.uniform(0.97, 1.03), 2)
            else:
                mdesc, cost, uom = None, 0.0, None
            rows.append({"Order": int(ev.order), "Posting Date": second_date if (second_date and i == len(lines) - 1) else ev.date,
                         "Work Description": description, "Material description": mdesc, "Cost (In Rs)": cost,
                         "Quantity": -q if code else 0, "Base Unit of Measure": uom, "G/L Account": 40000500,
                         "Cost Center": None, "Planning Department": dept, "Work Type": work_type,
                         "Movement type": 261, "Material": int(code) if code else None, "Storage location": "SN01", **header})
        truth.append({"order": ev.order, "asset_id": ev.asset_id, "kind": ev.kind, "component": ev.component or "",
                      "failure_mode": ev.failure_mode or "", "root_cause": ev.root_cause or "", "action": ev.action or "",
                      "symptom": ev.symptom or "", "breakdown": ev.breakdown})

    _write_xlsx(out_xlsx, rows)
    truth_csv.parent.mkdir(parents=True, exist_ok=True)
    with open(truth_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["order", "asset_id", "kind", "component", "failure_mode", "root_cause",
                                          "action", "symptom", "breakdown"], extrasaction="ignore")
        w.writeheader()
        w.writerows(truth)
    return {"rows": len(rows), "orders": sum(1 for t in truth if t["order"]), "file": str(out_xlsx)}


def _write_xlsx(path: Path, rows: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "SAP PM export"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for r in rows:
        ws.append([r.get(c) for c in COLUMNS])
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "m/d/yyyy"
        row[4].number_format = '"₹" #,##0.00'
    widths = {"A": 13, "B": 12, "C": 40, "D": 40, "E": 14, "O": 18, "P": 16, "Q": 22, "U": 10, "V": 40}
    for col, wdt in widths.items():
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A2"
    wb.save(path)
